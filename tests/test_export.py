from io import BytesIO
import pandas as pd
import pytest
from openpyxl import load_workbook
from axioma.demo import make_demo
from axioma.engine import calculate
from axioma.export import approve_orders, export_csv, export_xlsx, state_fingerprint
from axioma.importers import csv_bundle, read_canonical
import zipfile


def test_approval_export_roundtrip_and_protection():
    results, _ = calculate(make_demo())
    edited = results.assign(order_qty=results.recommended_qty)
    with pytest.raises(ValueError):
        approve_orders(results, edited, '', '2026-09-22')
    approved = approve_orders(results, edited, 'Тест', '2026-09-22')
    approved.loc[approved.index[0], 'name'] = '=HYPERLINK("bad")'
    csv = pd.read_csv(BytesIO(export_csv(approved)))
    assert len(csv) == len(approved)
    assert csv.iloc[0]['name'].startswith("'=")
    wb = load_workbook(BytesIO(export_xlsx(approved)))
    assert wb.active.max_row == len(approved) + 1
    assert all(c.data_type != 'f' for row in wb.active for c in row)


def test_approval_validates_quantity_and_moq():
    results, _ = calculate(make_demo())
    edited = results.assign(order_qty=results.recommended_qty)
    edited.loc[0, 'order_qty'] = -1
    with pytest.raises(ValueError):
        approve_orders(results, edited, 'Тест', '2026-09-22')
    edited.loc[0, 'order_qty'] = 1
    with pytest.raises(ValueError):
        approve_orders(results, edited, 'Тест', '2026-09-22')


def test_fingerprint_invalidates_changed_order():
    frame = pd.DataFrame({'order_qty': [10]})
    before = state_fingerprint(frame, {'reviewer': 'A'})
    frame.loc[0, 'order_qty'] = 20
    assert before != state_fingerprint(frame, {'reviewer': 'A'})


def test_csv_demo_roundtrip():
    original = make_demo()
    with zipfile.ZipFile(BytesIO(csv_bundle(original))) as archive:
        imported = read_canonical({n: archive.read(n) for n in archive.namelist()})
    assert len(imported.sales) == len(original.sales)
    assert len(imported.stockouts) == 1
    before, _ = calculate(original)
    after, _ = calculate(imported)
    assert before.recommended_qty.tolist() == after.recommended_qty.tolist()
