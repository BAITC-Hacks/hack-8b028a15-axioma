from copy import deepcopy
from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path
import os
import numpy as np
import pandas as pd
import streamlit as st
from src.data import demo_data, parse_files
from src.engine import Settings, calculate
from src.scenarios import stock_scenario, annotate_scenario
from src.orders import add_reviewed_lines

st.set_page_config(page_title='Axioma · Закупки',page_icon='◈',layout='wide')
st.markdown('''<style>
.block-container{padding-top:2rem;max-width:1450px} h1{letter-spacing:-1.5px}
[data-testid="stMetric"]{background:white;border:1px solid #e3e8ef;border-radius:14px;padding:18px}
[data-testid="stSidebar"]{border-right:1px solid #e3e8ef}
.eyebrow{font-size:12px;font-weight:700;letter-spacing:3px;color:#0d9488}
.intro{font-size:18px;color:#61718a;max-width:850px;margin-bottom:24px}
</style>''',unsafe_allow_html=True)

@st.cache_data(show_spinner=False)
def load_demo():return demo_data()
@st.cache_data(show_spinner=False)
def load_files(files,supplier):return parse_files(files,supplier)
@st.cache_data(show_spinner=False)
def run(ds,cfg):return calculate(ds,cfg)

LABELS={'stock_threshold':'Порог пополнения','stock_basis':'Основание остатка','snapshot_date':'Дата снимка','snapshot_stock':'Остаток в снимке','recorded_sales':'Продажи после снимка','scenario_stock':'Сценарный остаток','snapshot_age_days':'Возраст снимка, дней','sku':'Код 1С','article':'Артикул','name':'Товар','category':'Категория','supplier':'Поставщик','unit':'Ед.','stock':'Свободный остаток','in_transit':'Приедет в период','late_transit':'Приедет позже','forecast':'Прогноз спроса','safety':'Страховой запас','daily':'Спрос в день','growth':'Тренд, %','season':'Сезонный коэффициент','excluded':'Исключено всплесков','lost':'Восстановлено спроса','pack':'Кратность','moq':'Минимальная партия','recommended':'Рекомендация','status':'Статус','reason':'Обоснование','stockout_mode':'Оценка отсутствия'}
LABELS.update(method='Выбранный метод',demand_type='Характер спроса',cv_months='Месяцев внутренней проверки',cv_wape='Внутренняя ошибка, %',stress_recommended='Заказ в стресс-сценарии',first_shortage='Первый риск дефицита',source_growth_percent='Рост из сводки, %',quantity='К заказу',calculated_at='Дата расчёта',planning_method='Метод расчёта')

def export_excel(frame, cfg):
    buffer=BytesIO()
    safe=frame.rename(columns=LABELS).copy()
    # Prevent spreadsheet formula injection through imported product names.
    for col in safe.select_dtypes(include=['object','str']).columns:
        safe[col]=safe[col].map(lambda x:"'"+x if isinstance(x,str) and x.startswith(('=','+','-','@')) else x)
    with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
        safe.to_excel(writer,index=False,sheet_name='Рекомендации')
        parameters=cfg if isinstance(cfg,dict) else vars(cfg)
        pd.DataFrame([{'Параметр':k,'Значение':str(v)} for k,v in parameters.items()]).to_excel(writer,index=False,sheet_name='Параметры')
        ws=writer.sheets['Рекомендации'];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='0D9488')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(60,max(14,len(str(col[0].value))+3))
    return buffer.getvalue()

with st.sidebar:
    st.markdown('## ◈ AXIOMA')
    st.caption('Помощник менеджера закупок')
    mode=st.radio('Источник данных',['Показать пример','Файлы поставщика'])
    supplier=st.selectbox('Поставщик',['Systeme Electric','IEK']) if mode=='Файлы поставщика' else 'Демо'
    st.divider()
    st.markdown('**Период планирования**')
    as_of=st.date_input('Дата расчёта',date(2026,9,23))
    lead=st.number_input('Срок новой поставки, дней',min_value=1,max_value=365,value=21,help='Допущение. Уточните срок у поставщика.')
    review=st.number_input('До следующего заказа, дней',min_value=1,max_value=90,value=14)
    safety=st.number_input('Страховой запас, дней',min_value=0,max_value=90,value=7)
    growth=st.slider('Дополнительный прогноз прироста, %',-50,100,0,help='Ручной сценарий поверх тренда, оценённого по истории.')
    st.divider()
    method=st.selectbox('Метод расчёта',['Общая ML-модель','Автовыбор по истории','Axioma: настроенная модель','Обычное среднее за 6 месяцев'],help='Общая модель обучается на прошлых месяцах всех загруженных товаров. При менее 200 обучающих примерах применяется статистический автовыбор. Автовыбор сравнивает методы на шести доступных месяцах внутри истории. Независимая проверка выполняется отдельно.')
    source_growth=st.checkbox('Применить рост из сводки вместо тренда',False,help='Поле «Кэф. Роста» импортируется как доля изменения: 0,2 означает +20%. Это явный сценарий менеджера; подтвердите смысл коэффициента у владельца данных.')
    remove=st.toggle('Исключать разовые всплески',True)
    seasonal=st.toggle('Учитывать сезонность',True)
    trend=st.toggle('Учитывать устойчивый рост',True)
    compensate=st.toggle('Восстанавливать упущенный спрос',True)
    approx=st.checkbox('Приближение по месячным остаткам',False,help='Только явный нулевой снимок. Это не точные даты отсутствия товара.')
    st.caption('Расчёт выполняется локально. API-ключи и платные подписки не нужны.')

st.markdown('<div class="eyebrow">AXIOMA / INVENTORY INTELLIGENCE</div>',unsafe_allow_html=True)
st.title('Заказывайте то, что нужно.')
st.markdown('<div class="intro">Продажи, склад и ожидаемые поставки — в одном расчёте. Каждая рекомендация сопровождается понятным объяснением.</div>',unsafe_allow_html=True)

if mode=='Показать пример':
    ds=deepcopy(load_demo())
    st.info('Демонстрационный набор: 6 товаров, сезонность, рост, разовая покупка на 1 500 единиц и известный период отсутствия товара. Это синтетические данные.')
else:
    uploaded=st.file_uploader('Загрузите все 6 Excel-файлов выбранного поставщика',type=['xlsx'],accept_multiple_files=True)
    local_root=os.environ.get('AXIOMA_DATA_DIR')
    local_files=[]
    if local_root:
        folder=Path(local_root)/('Systeme electric' if supplier=='Systeme Electric' else 'IEK')
        if folder.is_dir():
            if st.button('Открыть предоставленные локальные файлы',type='primary'):st.session_state['local_supplier']=supplier
            if st.session_state.get('local_supplier')==supplier:local_files=[(p.name,p.read_bytes()) for p in sorted(folder.glob('*.xlsx'))]
    inputs=[(f.name,f.getvalue()) for f in uploaded] or local_files
    if not inputs:
        st.markdown('**Начните с примера слева или выберите файлы выше.** После загрузки здесь появятся рекомендации.')
        st.stop()
    with st.spinner('Читаю файлы и связываю товары по коду 1С…'):ds=deepcopy(load_files(inputs,supplier))
    if ds.products.empty:st.error('Не удалось найти товары. Проверьте отчёт об импорте ниже.');st.write(ds.warnings);st.stop()

with st.expander('Данные и допущения · проверить перед заказом',expanded=False):
    st.dataframe(pd.DataFrame(ds.sources),hide_index=True,width='stretch')
    for warning in ds.warnings:st.warning(warning)
    st.caption('Для расчёта используются завершённые месяцы. Отрицательные продажи сохраняются как корректировки; отрицательный спрос не прогнозируется. Срок поставки и страховой запас задаются пользователем.')
    missing=ds.products.stock.isna().sum()
    st.markdown(f'**Актуальный остаток не указан: {missing} позиций.** Можно исправить таблицу ниже или загрузить CSV.')
    st.caption('CSV остатков: sku,stock,category,pack,moq. Обязательны sku и stock; коды храните текстом. pack — кратность в единицах учёта. Для бухт переведите кратность в метры.')
    overrides=st.file_uploader('Актуальные остатки и параметры товаров',type=['csv'])
    if overrides:
        try:
            extra=pd.read_csv(overrides,dtype={'sku':str})
            if not {'sku','stock'}.issubset(extra.columns):raise ValueError('Нужны колонки sku и stock')
            if extra.sku.duplicated().any():raise ValueError('Повторяются коды SKU')
            extra=extra.set_index('sku');base=ds.products.set_index('sku')
            for c in ['stock','pack','moq']:
                if c in extra:extra[c]=pd.to_numeric(extra[c],errors='raise')
            if 'stock' in extra and (extra.stock<0).any():raise ValueError('Свободный остаток должен быть неотрицательным')
            if any(c in extra and (extra[c]<=0).any() for c in ['pack','moq']):raise ValueError('Кратность и минимум должны быть положительными')
            base.update(extra[[c for c in ['stock','category','pack','moq'] if c in extra]])
            ds.products=base.reset_index();st.caption(f'Обновлено {len(base.index.intersection(extra.index))} товаров; не найдено {len(extra.index.difference(base.index))}.')
        except Exception as exc:st.error(f'CSV не применён: {exc}')
    changed=st.data_editor(ds.products.rename(columns=LABELS),disabled=['Код 1С','Артикул','Товар','Ед.'],hide_index=True,width='stretch',key=f'products-{mode}-{supplier}',column_config={'Свободный остаток':st.column_config.NumberColumn(min_value=0),'Кратность':st.column_config.NumberColumn(min_value=1),'Минимальная партия':st.column_config.NumberColumn(min_value=1)})
    ds.products=changed.rename(columns={v:k for k,v in LABELS.items()})
    st.caption('CSV отсутствия товара: sku,start,end, даты YYYY-MM-DD. Оба конца периода включены.')
    stockout_file=st.file_uploader('Подтверждённые периоды отсутствия',type=['csv'])
    if stockout_file:
        try:
            so=pd.read_csv(stockout_file,dtype={'sku':str})
            if not {'sku','start','end'}.issubset(so.columns):raise ValueError('Нужны sku,start,end')
            for c in ['start','end']:so[c]=pd.to_datetime(so[c],errors='raise')
            if (so.end<so.start).any() or so[['start','end']].isna().any().any():raise ValueError('Некорректные границы периода')
            ds.stockouts=so
        except Exception as exc:st.error(f'Периоды не применены: {exc}')
    st.markdown('**Правила категорий**')
    st.caption('Множитель страхового запаса. По умолчанию 1 для всех: бизнес-значение кодов категорий не предоставлено. Установите согласованные с менеджером правила.')
    categories=sorted(ds.products.category.astype(str).unique())
    cat_frame=st.data_editor(pd.DataFrame({'Категория':categories,'Множитель':[1.]*len(categories)}),disabled=['Категория'],hide_index=True,key=f'cat-{mode}-{supplier}',column_config={'Множитель':st.column_config.NumberColumn(min_value=0.,max_value=5.)})
    category_factors=dict(zip(cat_frame['Категория'],cat_frame['Множитель']))

stock_evidence=pd.DataFrame()
if ds.products.stock.isna().any() and not ds.stocks.empty:
    st.markdown('**Расчёт при неполных данных об остатках**')
    scenario=st.radio('Как рассчитать неизвестные остатки?', ['Только известные остатки','Сценарий: остаток на начало месяца','Сценарий: вычесть продажи без поступлений'], horizontal=True)
    if scenario!='Только известные остатки':
        ds,stock_evidence=stock_scenario(ds,str(as_of),'reference' if scenario=='Сценарий: остаток на начало месяца' else 'depletion')
        st.warning(f'Сценарный расчёт для {len(stock_evidence)} товаров. Поступления, резервы и другие движения после снимка неизвестны. Это не фактические остатки и не границы возможного запаса. Позиции без достаточных данных остаются неизвестными.')
        with st.expander('Из чего получились сценарные остатки'):
            st.dataframe(stock_evidence.rename(columns=LABELS),hide_index=True,width='stretch')

cfg=Settings(str(as_of),int(lead),int(review),int(safety),float(growth),remove,seasonal,trend,compensate,approx,category_factors)
cfg=replace(cfg,forecast_method={'Автовыбор по истории':'adaptive','Общая ML-модель':'pooled'}.get(method,'legacy'),use_source_growth=source_growth)
if method=='Обычное среднее за 6 месяцев':
    cfg=replace(cfg,remove_outliers=False,seasonality=False,trend=False,compensate_stockout=False)
    st.info('Выбрано обычное среднее. Переключатели очистки, сезонности, тренда и восстановления спроса не применяются. Ручной прирост и параметры запаса сохранены.')
with st.spinner('Считаю спрос и заказы…'):result,histories,anomalies=run(ds,cfg)
if not stock_evidence.empty:result=annotate_scenario(result,stock_evidence)
if result.empty:st.warning('Нет позиций для расчёта');st.stop()
result=result.assign(_priority=result.status.map({'Срочно':0,'Заказать':1,'Нужен остаток':2,'Нет истории':3,'Достаточно':4})).sort_values(['_priority','sku']).drop(columns='_priority').reset_index(drop=True)
cards=st.columns(4)
cards[0].metric('Товаров в анализе',len(result))
cards[1].metric('Нужно заказать',int(result.recommended.gt(0).sum()))
cards[2].metric('Риск дефицита',int(result.status.eq('Срочно').sum()))
cards[3].metric('Нужны данные',int(result.recommended.isna().sum()))
orders_tab,detail_tab,quality_tab,comparison_tab=st.tabs(['Рекомендации','Почему столько','Качество данных','Сравнение методов'])
with orders_tab:
    a,b=st.columns([2,1]);search=a.text_input('Найти товар',placeholder='Название, артикул или код')
    statuses=b.multiselect('Статус',list(result.status.unique()),default=list(result.status.unique()))
    selected=result[result.status.isin(statuses)].copy()
    if search:selected=selected[selected[['sku','article','name']].astype(str).apply(lambda s:s.str.contains(search,case=False,regex=False)).any(axis=1)]
    chosen_categories=st.multiselect('Категории',categories)
    if chosen_categories:selected=selected[selected.category.astype(str).isin(chosen_categories)]
    columns=['stock_basis','status','first_shortage','sku','article','name','unit','stock','stock_threshold','in_transit','forecast','recommended','method','demand_type']
    st.dataframe(selected[[c for c in columns if c in selected]].rename(columns=LABELS),hide_index=True,width='stretch')
    st.download_button('Скачать все рекомендации · Excel',export_excel(result,cfg),file_name=f'axioma-{as_of}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with st.expander('Проверить и утвердить заказ'):
        draft=result[result.recommended.gt(0)].copy()
        if draft.empty:st.info('Нет положительных рекомендаций для утверждения.')
        else:
            draft=draft[[c for c in ['sku','article','name','supplier','recommended','pack','moq','stock_basis','snapshot_date'] if c in draft]].copy();draft['Утвердить']=False;draft['К заказу']=draft.recommended
            reviewed=st.data_editor(draft.rename(columns=LABELS),hide_index=True,disabled=['Код 1С','Артикул','Товар','Поставщик','Рекомендация','Кратность','Минимальная партия','Основание остатка','Дата снимка'],column_config={'К заказу':st.column_config.NumberColumn(min_value=0)},key=f'approval-{mode}-{supplier}')
            approved=reviewed[reviewed['Утвердить']&reviewed['К заказу'].gt(0)]
            bad=approved[(approved['К заказу']<approved['Минимальная партия'])|~np.isclose(approved['К заказу']%approved['Кратность'],0)]
            if len(bad):st.error('Исправьте количество: оно должно соответствовать минимуму и кратности.')
            acknowledge=True if stock_evidence.empty else st.checkbox('Понимаю: это сценарный заказ с предполагаемыми остатками')
            st.download_button('Скачать утверждённые позиции',export_excel(approved,cfg),file_name='axioma-approved.xlsx',disabled=approved.empty or not bad.empty or not acknowledge)
            if st.button('Добавить утверждённые позиции в общую корзину',disabled=approved.empty or not bad.empty or not acknowledge):
                lines=approved.rename(columns={v:k for k,v in LABELS.items()}).copy()
                lines=lines.drop(columns=['Утвердить'],errors='ignore')
                lines=lines.merge(result[['supplier','sku','reason']],on=['supplier','sku'],validate='one_to_one')
                lines['calculated_at']=str(as_of);lines['planning_method']=method
                st.session_state['order_cart']=add_reviewed_lines(st.session_state.get('order_cart',pd.DataFrame()),lines)
                st.success(f'Добавлено {len(lines)} позиций. Теперь можно рассчитать другого поставщика.')
            st.caption('Скачивание не отправляет заказ поставщику. Утверждение действует в текущей сессии; сохраните файл.')
    cart=st.session_state.get('order_cart',pd.DataFrame())
    if not cart.empty:
        st.subheader('Общая корзина поставщиков')
        st.caption('Сохранённые утверждения этой сессии. Новые расчёты не меняют их автоматически: повторное добавление заменяет строку того же поставщика и товара. Корзина хранит дату и метод каждого расчёта.')
        for supplier_name,group in cart.groupby('supplier'):
            with st.expander(f'{supplier_name} · {len(group)} позиций',expanded=True):
                st.dataframe(group.rename(columns=LABELS),hide_index=True,width='stretch')
        st.download_button('Скачать общую корзину · Excel',export_excel(cart,{}),file_name='axioma-supplier-orders.xlsx')
        if st.button('Очистить корзину этой сессии'):
            st.session_state['order_cart']=pd.DataFrame();st.rerun()
with detail_tab:
    sku=st.selectbox('Выберите товар',list(histories),format_func=lambda k:f"{k} · {result.set_index('sku').loc[k,'name']}") if histories else None
    if sku:
        row=result.set_index('sku').loc[sku];st.subheader(row['name']);st.write(row.reason)
        c1,c2,c3=st.columns(3);c1.metric('Рекомендовано',f'{row.recommended:g}' if pd.notna(row.recommended) else 'Нужен остаток');c2.metric('Порог пополнения',f'{row.stock_threshold:g}');c3.metric('Восстановлено спроса',row.lost)
        st.write('Метод:',row.method,'· Характер спроса:',row.demand_type)
        st.line_chart(histories[sku].set_index('Дата'),color=['#94A3B8','#0D9488','#F59E0B'])
        st.caption('Серый — фактические продажи; зелёный — без разовых всплесков; оранжевый — с компенсацией отсутствия товара.')
        st.write('Оценка отсутствия:',row.stockout_mode)
        model_info=histories[sku].attrs.get('model_info',{})
        if model_info.get('candidates'):
            with st.expander('Почему выбран этот метод'):
                st.write(f'Сравнение на {row.cv_months:g} последних доступных месяцах внутри истории. Это выбор метода, а не независимая оценка качества. При близкой ошибке усредняются до трёх методов.')
                st.dataframe(pd.DataFrame(model_info['candidates']).rename(columns={'method':'Метод','mae_daily':'Средняя абсолютная ошибка, ед./день'}),hide_index=True,width='stretch')
        future_frame=pd.DataFrame(histories[sku].attrs.get('future',[]))
        if not future_frame.empty and pd.notna(row.stock):
            future_frame['Дата']=pd.to_datetime(future_frame['Дата'])
            transit=ds.transit[ds.transit.sku.eq(sku)]
            arrivals=transit.groupby('eta').quantity.sum() if not transit.empty else pd.Series(dtype=float)
            future_frame['Ожидаемые поступления']=future_frame['Дата'].map(arrivals).fillna(0).clip(lower=0)
            future_frame['Без нового заказа']=row.stock+(future_frame['Ожидаемые поступления']-future_frame['Прогноз в день']).cumsum()
            future_frame['С новым заказом']=future_frame['Без нового заказа']+np.where(future_frame['Дата']>=pd.Timestamp(as_of)+pd.Timedelta(days=lead),row.recommended,0)
            st.markdown('**Что будет с запасом по дням**')
            st.line_chart(future_frame.set_index('Дата')[['Без нового заказа','С новым заказом']],color=['#EF4444','#0D9488'])
            st.caption('Отрицательный баланс показывает неудовлетворённую потребность. Поступления учитываются в начале дня; новый заказ — через заданный срок. Это сценарий по прогнозу, а не фактическое движение склада.')
            if pd.notna(row.first_shortage):st.warning(f'Риск дефицита с {row.first_shortage}, до прихода нового заказа. Нужна ускоренная поставка или перемещение со склада.')
            if cfg.forecast_method=='adaptive' and row.cv_months>=3:
                st.info(f'Стресс-сценарий: {row.stress_recommended:g} ед. к заказу, если спрос превысит прогноз на величину 90-го процентиля прошлых положительных ошибок. Это ориентир для проверки, не гарантированный уровень сервиса.')
with quality_tab:
    st.subheader('Прозрачный расчёт')
    st.write('Заказ = прогноз на срок поставки и интервал пересмотра + страховой запас − свободный остаток − поступления в этот период. Положительный результат округляется с учётом минимума и кратности.')
    st.write('Крупные накладные и всплески клиента за день проверяются устойчивым порогом по распределению объёмов. Подозрительный объём заменяется типичным; исходные файлы не изменяются.')
    st.dataframe(pd.DataFrame(ds.sources),hide_index=True,width='stretch')
    if not anomalies.empty:st.markdown('**Обнаруженные всплески**');st.dataframe(anomalies,hide_index=True,width='stretch')
    else:st.info('Всплески не обнаружены либо нет детальных продаж для их проверки.')
    st.warning('Без выбранного сценария неизвестный остаток блокирует заказ. Сценарные остатки всегда помечены в результатах. Неизвестные или просроченные даты поступления требуют уточнения. Без client_id невозможно проверить покупки одного клиента между разными накладными.')
    st.caption('Прототип HackAlem AI · AI-агент Codex использован при разработке. Расчёт детерминированный; внешние AI API не вызываются.')

with comparison_tab:
    st.subheader('Что меняет обработка спроса')
    st.write('Сравните три расчёта на одинаковых товарах, остатках и сроках: обычное среднее за 6 завершённых месяцев; текущая модель без удаления всплесков; выбранная вами модель.')
    st.caption('Это сравнение рекомендаций, а не проверка точности на будущих продажах. Меньший заказ сам по себе не доказывает экономию. Количества разных товаров не складываются.')
    if st.button('Рассчитать сравнение'):
        with st.spinner('Сравниваю методы на одинаковых входных данных…'):
            simple=run(ds,replace(cfg,forecast_method='legacy',remove_outliers=False,seasonality=False,trend=False,compensate_stockout=False))[0]
            unclean=run(ds,replace(cfg,remove_outliers=False))[0]
        comparison=result[[c for c in ['sku','name','recommended','stock_basis','snapshot_date'] if c in result]].rename(columns={'recommended':'Выбранная модель'})
        comparison=comparison.merge(simple[['sku','recommended']].rename(columns={'recommended':'Обычное среднее'}),on='sku',validate='one_to_one')
        comparison=comparison.merge(unclean[['sku','recommended']].rename(columns={'recommended':'Без удаления всплесков'}),on='sku',validate='one_to_one')
        comparison['Разница со средним']=comparison['Выбранная модель']-comparison['Обычное среднее']
        comparison['Влияние фильтра всплесков']=comparison['Выбранная модель']-comparison['Без удаления всплесков']
        comparison=comparison.sort_values('Влияние фильтра всплесков',key=lambda x:x.abs(),ascending=False)
        st.dataframe(comparison.rename(columns=LABELS),hide_index=True,width='stretch')
        st.download_button('Скачать сравнение методов',export_excel(comparison,cfg),file_name='axioma-method-comparison.xlsx')

    st.divider()
    st.subheader('Проверка на прошлых месяцах')
    st.write('Предсказываем 3 завершённых месяца по очереди: каждый прогноз видит только более раннюю историю. Затем сравниваем его с фактическими продажами и обычным средним.')
    st.caption('Для быстрого запуска: первые 30 кодов с минимум 6 месяцами истории до проверки. Это выборка, не весь каталог. Готовые сезонные коэффициенты и периоды отсутствия исключены, поскольку неизвестно, когда они стали доступны. Факт продаж не равен скрытому спросу при дефиците и включает разовые покупки.')
    if st.button('Проверить прогноз на истории'):
        from src.validation import backtest
        with st.spinner('Проверяю прогноз без доступа к будущим продажам…'):
            detail,scores,coverage=backtest(ds,str(as_of))
        st.write(f"Проверены {coverage['evaluated_skus']} товаров из {coverage['eligible_skus']} подходящих. Месяцы: {', '.join(coverage['months'])}.")
        valid=scores.dropna(subset=['model_wape','baseline_wape'])
        if valid.empty:st.info('Недостаточно наблюдаемых продаж для оценки ошибки.')
        else:
            c1,c2,c3=st.columns(3)
            c1.metric('Средняя ошибка Axioma по товарам',f'{valid.model_wape.mean():.1f}%')
            c2.metric('Средняя ошибка обычного среднего',f'{valid.baseline_wape.mean():.1f}%')
            c3.metric('Средняя ошибка автовыбора',f'{valid.adaptive_wape.mean():.1f}%')
            st.caption('Для каждого товара WAPE = сумма абсолютных ошибок / сумма факта; затем берём среднее по товарам с положительным фактом. Меньше — лучше, ошибка может быть выше 100%. Это не «процент точности» и не денежная экономия.')
            st.dataframe(scores.drop(columns=['pooled_wape'],errors='ignore').rename(columns={'sku':'Код 1С','months':'Месяцев','actual_total':'Факт за период','model_wape':'Ошибка исходной модели, %','baseline_wape':'Ошибка среднего, %','adaptive_wape':'Ошибка автовыбора, %'}),hide_index=True,width='stretch')
            st.download_button('Скачать результаты проверки',export_excel(detail.rename(columns={'month':'Месяц','actual':'Факт','model':'Прогноз Axioma','baseline':'Прогноз среднего'}),cfg),file_name='axioma-validation.xlsx')
