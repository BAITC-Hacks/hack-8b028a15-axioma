from io import BytesIO
from datetime import datetime, timezone
import hashlib
import json
import math
import pandas as pd


def safe_text(value):
    # Prevent spreadsheet formula injection from imported product names.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def approve_orders(recommendations, edited, reviewer, as_of):
    if not reviewer.strip():
        raise ValueError('Укажите имя или инициалы ответственного.')
    merged = recommendations.merge(edited[['supplier', 'sku', 'order_qty']], on=['supplier', 'sku'], how='inner', validate='one_to_one')
    if len(merged) != len(edited):
        raise ValueError('В заказе есть позиции без рекомендации.')
    qty = pd.to_numeric(merged.order_qty, errors='coerce')
    if qty.isna().any() or (qty < 0).any():
        raise ValueError('Количество заказа должно быть неотрицательным числом.')
    if ((merged.recommended_qty.isna()) & (qty > 0)).any():
        raise ValueError('Нельзя утвердить заказ по позициям с неполными данными.')
    for row in merged[qty > 0].itertuples():
        if row.order_qty < row.moq or not math.isclose(row.order_qty / row.pack, round(row.order_qty / row.pack), abs_tol=1e-8):
            raise ValueError(f'{row.sku}: нарушены минимальная партия или кратность.')
    merged['approved_by'] = reviewer.strip()
    merged['approved_at_utc'] = datetime.now(timezone.utc).isoformat()
    merged['calculation_date'] = str(as_of)
    merged['approved_stock_units'] = qty * merged.purchase_factor
    return merged[qty > 0].copy()


def export_csv(frame):
    return frame.map(safe_text).to_csv(index=False).encode('utf-8-sig')


def export_xlsx(frame):
    # Application-generated order export, not a hand-authored spreadsheet artifact.
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        frame.map(safe_text).to_excel(writer, sheet_name='Утверждённый заказ', index=False)
        ws = writer.sheets['Утверждённый заказ']
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='173D38')
        for column in ws.columns:
            width = min(55, max(14, max(len(str(c.value or '')) for c in list(column)[:100]) + 2))
            ws.column_dimensions[column[0].column_letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    return buffer.getvalue()


def state_fingerprint(frame, settings):
    value = frame.to_json(date_format='iso') + json.dumps(settings, sort_keys=True, default=str)
    return hashlib.sha256(value.encode()).hexdigest()
